from __future__ import annotations

import io
import zipfile
from pathlib import Path

import fitz
import pytest

from app.core.errors import (
    FileTooLargeError,
    NoExtractableTextError,
    TooManyPagesError,
    UnreadableDocumentError,
    UnsupportedFileTypeError,
)
from app.services.chunking import chunk_segments
from app.services.file_validation import FileType, check_size, detect_file_type, sanitize_filename
from app.services.parsing import parse_document

SAMPLES = Path(__file__).parent / "samples"


def _read(name: str) -> bytes:
    return (SAMPLES / name).read_bytes()


# ---------- upload checks ----------


def test_detects_real_pdf_and_docx() -> None:
    assert detect_file_type(_read("employee_handbook.pdf"), "employee_handbook.pdf") == FileType.PDF
    assert detect_file_type(_read("cloud_service_agreement.docx"), "a.DOCX") == FileType.DOCX


def test_rejects_unsupported_extension() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(b"hello", "notes.txt")


def test_rejects_fake_pdf_with_pdf_extension() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(b"<html>not a pdf</html>", "evil.pdf")


def test_rejects_pdf_renamed_to_docx_and_back() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(_read("employee_handbook.pdf"), "x.docx")
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(_read("cloud_service_agreement.docx"), "x.pdf")


def test_rejects_generic_zip_and_xlsx_like_archive_named_docx() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")  # an xlsx, not a docx
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(buffer.getvalue(), "sheet.docx")


def test_rejects_empty_file() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(b"", "a.pdf")


def test_size_limit() -> None:
    check_size(10, 10)
    with pytest.raises(FileTooLargeError):
        check_size(11, 10)


def test_sanitize_filename_strips_paths_and_control_chars() -> None:
    assert sanitize_filename("..\\..\\etc/passwd") == "passwd"
    assert sanitize_filename("a\x00b\nc.pdf") == "abc.pdf"
    assert sanitize_filename("") == "document"


# ---------- parsing ----------


def test_pdf_text_per_page_with_page_numbers() -> None:
    parsed = parse_document(_read("employee_handbook.pdf"), FileType.PDF, max_pages=50)
    assert parsed.page_count == 6
    assert [s.page for s in parsed.segments] == [1, 2, 3, 4, 5, 6]
    assert "25 days of paid annual leave" in parsed.segments[1].text
    assert "\n" not in parsed.segments[1].text  # whitespace normalized


def test_pdf_page_limit() -> None:
    with pytest.raises(TooManyPagesError):
        parse_document(_read("employee_handbook.pdf"), FileType.PDF, max_pages=3)


def test_pdf_without_text_layer_is_rejected() -> None:
    doc = fitz.open()
    page = doc.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 50, 50), False)
    page.insert_image(fitz.Rect(10, 10, 200, 200), pixmap=pixmap)  # image only, like a scan
    data = doc.tobytes()
    with pytest.raises(NoExtractableTextError):
        parse_document(data, FileType.PDF, max_pages=10)


def test_corrupt_pdf_is_unreadable() -> None:
    with pytest.raises(UnreadableDocumentError):
        parse_document(b"%PDF-1.7\nthis is garbage", FileType.PDF, max_pages=10)


def test_corrupt_docx_is_unreadable() -> None:
    with pytest.raises(UnreadableDocumentError):
        parse_document(b"PK\x03\x04garbage", FileType.DOCX, max_pages=10)


def test_docx_has_sections_without_pages() -> None:
    parsed = parse_document(_read("cloud_service_agreement.docx"), FileType.DOCX, max_pages=50)
    assert parsed.page_count is None
    assert all(s.page is None for s in parsed.segments)
    sections = [s.section for s in parsed.segments]
    assert "Helios Cloud Service Agreement > 1. Service Levels" in sections
    assert "Helios Cloud Service Agreement > 1. Service Levels > 1.1 Service Credits" in sections
    # table rows are flattened into text
    support = next(s for s in parsed.segments if s.section and s.section.endswith("2. Support"))
    assert "P1 - service down | 1 hour | 24/7 on the Premium plan only" in support.text


def _docx_bytes(build) -> bytes:  # type: ignore[no-untyped-def]
    from docx import Document

    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_vertically_merged_cell_is_extracted_once() -> None:
    # Regression: a CV laid out as a table with one tall merged column was indexed twice.
    def build(document) -> None:  # type: ignore[no-untyped-def]
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(1, 0).text = "Contact"
        merged = table.cell(0, 1).merge(table.cell(1, 1))
        merged.text = "Long profile text that spans both rows"

    text = parse_document(_docx_bytes(build), FileType.DOCX, max_pages=10).segments[0].text
    assert text.count("Long profile text that spans both rows") == 1


def test_docx_horizontally_merged_cell_once_but_equal_neighbours_kept() -> None:
    def build(document) -> None:  # type: ignore[no-untyped-def]
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "Title across all columns"
        for column in range(3):
            table.cell(1, column).text = "Yes"  # genuinely identical, separate cells

    text = parse_document(_docx_bytes(build), FileType.DOCX, max_pages=10).segments[0].text
    assert text.count("Title across all columns") == 1
    assert "Yes | Yes | Yes" in text


def test_sample_documents_chunk_cleanly() -> None:
    for name, file_type in [
        ("employee_handbook.pdf", FileType.PDF),
        ("cloud_service_agreement.docx", FileType.DOCX),
        ("facilities_report_injection.pdf", FileType.PDF),
        ("router_manual.pdf", FileType.PDF),
    ]:
        parsed = parse_document(_read(name), file_type, max_pages=50)
        chunks = chunk_segments("d", name, parsed.segments, 500, 100)
        assert chunks and all(0 < len(c.text) <= 500 for c in chunks)


def test_hidden_injection_text_is_extracted_as_plain_text() -> None:
    """Documents the threat: invisible white text is in the extracted text, so we must treat all
    document text as untrusted (see answering.py for the actual defenses)."""
    parsed = parse_document(_read("facilities_report_injection.pdf"), FileType.PDF, max_pages=10)
    assert "SYSTEM OVERRIDE" in parsed.segments[1].text
